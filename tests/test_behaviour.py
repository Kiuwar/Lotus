"""Testes de conteúdo: a saída precisa estar CERTA, não só existir."""

from __future__ import annotations

import csv
import json
from pathlib import Path

import pytest
from PIL import Image

from conversor.converters import media
from conversor.core import registry
from conversor.core.detect import DetectionError, detect
from conversor.core.engine import ConversionError, convert_file

from .conftest import HAS_FFMPEG


# ---------- dados ----------

def test_csv_to_xlsx_preserves_text_and_blocks_formulas(sample, tmp_path):
    from openpyxl import load_workbook

    out = convert_file(sample("csv"), "xlsx", out_dir=tmp_path).output
    ws = load_workbook(out).active
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    assert rows[0] == ["nome", "cep", "preço"]
    assert rows[1] == ["João", "01234-000", "10,50"]  # zero à esquerda e vírgula intactos
    assert ws["B3"].value == "=SOMA(1;2)"
    assert ws["B3"].data_type == "s", "texto com '=' não pode virar fórmula"


def test_json_to_csv_unions_keys_and_serialises_nested(sample, tmp_path):
    res = convert_file(sample("json"), "csv", out_dir=tmp_path)
    with open(res.output, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f, delimiter=";"))
    assert rows[0] == ["nome", "cep", "itens", "idade"]
    assert rows[1] == ["João", "01234-000", "[1, 2]", ""]
    assert rows[2] == ["Ana", "", "", "30"]
    assert any("aninhados" in w for w in res.warnings)


def test_xlsx_to_json_keeps_types(sample, tmp_path):
    out = convert_file(sample("xlsx"), "json", out_dir=tmp_path).output
    assert json.loads(out.read_text(encoding="utf-8")) == [{"nome": "Café", "valor": 3.5}]


def test_csv_cp1252_is_decoded_with_warning(tmp_path):
    src = tmp_path / "antigo.csv"
    src.write_bytes("nome;cidade\nJosé;São Paulo\n".encode("cp1252"))
    res = convert_file(src, "json", out_dir=tmp_path)
    assert json.loads(res.output.read_text(encoding="utf-8")) == [
        {"nome": "José", "cidade": "São Paulo"}
    ]
    assert any("Windows-1252" in w for w in res.warnings)


def test_json_object_root_is_rejected_clearly(tmp_path):
    src = tmp_path / "obj.json"
    src.write_text('{"a": 1, "b": 2}', encoding="utf-8")
    with pytest.raises(ConversionError, match="lista de objetos"):
        convert_file(src, "csv", out_dir=tmp_path)


# ---------- imagens ----------

def test_png_with_alpha_to_jpg_flattens_on_white(tmp_path):
    src = tmp_path / "t.png"
    Image.new("RGBA", (10, 10), (0, 0, 0, 0)).save(src)
    res = convert_file(src, "jpg", out_dir=tmp_path)
    with Image.open(res.output) as im:
        assert im.format == "JPEG"
        r, g, b = im.convert("RGB").getpixel((5, 5))
        assert min(r, g, b) > 245  # branco (JPEG tem pequena perda)
    assert any("transparência" in w for w in res.warnings)


def test_exif_orientation_is_applied(tmp_path):
    src = tmp_path / "foto.jpg"
    im = Image.new("RGB", (40, 20), "red")
    exif = Image.Exif()
    exif[0x0112] = 6  # rotacionar 90° (padrão de celular em retrato)
    im.save(src, exif=exif)
    out = convert_file(src, "png", out_dir=tmp_path).output
    with Image.open(out) as o:
        assert o.size == (20, 40)


def test_animated_gif_warns_first_frame_only(tmp_path):
    src = tmp_path / "anim.gif"
    frames = [Image.new("RGB", (8, 8), c) for c in ("red", "blue")]
    frames[0].save(src, save_all=True, append_images=frames[1:])
    res = convert_file(src, "png", out_dir=tmp_path)
    assert any("2 quadros" in w for w in res.warnings)


def test_ico_contains_multiple_sizes(tmp_path):
    src = tmp_path / "logo.png"
    Image.new("RGBA", (256, 256), "green").save(src)
    out = convert_file(src, "ico", out_dir=tmp_path).output
    with Image.open(out) as im:
        assert {(16, 16), (32, 32), (256, 256)} <= set(im.info["sizes"])


# ---------- documentos ----------

def test_txt_to_pdf_keeps_accents(sample, tmp_path):
    from pypdf import PdfReader

    out = convert_file(sample("txt"), "pdf", out_dir=tmp_path).output
    text = "".join(p.extract_text() for p in PdfReader(out).pages)
    assert "ação" in text and "coração" in text


def test_md_to_html_renders_table(sample, tmp_path):
    html = convert_file(sample("md"), "html", out_dir=tmp_path).output.read_text(encoding="utf-8")
    assert "<h1>Título</h1>" in html and "<table>" in html and 'charset="utf-8"' in html


# ---------- segurança do motor ----------

def test_never_overwrites_by_default(sample, tmp_path):
    src = sample("png")
    a = convert_file(src, "jpg", out_dir=tmp_path).output
    b = convert_file(src, "jpg", out_dir=tmp_path).output
    assert a.name == "amostra.jpg" and b.name == "amostra (1).jpg"


def test_overwrite_when_asked(sample, tmp_path):
    src = sample("png")
    a = convert_file(src, "jpg", out_dir=tmp_path).output
    b = convert_file(src, "jpg", out_dir=tmp_path, overwrite=True).output
    assert a == b


def test_wrong_extension_uses_real_content(tmp_path):
    src = tmp_path / "mentira.png"
    Image.new("RGB", (8, 8)).save(src, format="JPEG")
    res = convert_file(src, "webp", out_dir=tmp_path)
    assert res.source_format == "jpg"
    assert any("conteúdo é JPG" in w for w in res.warnings)


def test_csv_starting_with_bm_is_still_csv(tmp_path):
    src = tmp_path / "carros.csv"
    src.write_text("BMW;Audi\n1;2\n", encoding="utf-8")
    assert detect(src).fmt == "csv"


def test_binary_with_text_extension_is_rejected(tmp_path):
    src = tmp_path / "x.txt"
    src.write_bytes(b"\x00\x01\x02garbage")
    with pytest.raises(DetectionError):
        detect(src)


def test_empty_file_is_rejected(tmp_path):
    src = tmp_path / "vazio.png"
    src.write_bytes(b"")
    with pytest.raises(ConversionError, match="vazio"):
        convert_file(src, "jpg")


def test_unsupported_pair_lists_alternatives(sample, tmp_path):
    with pytest.raises(ConversionError, match="não suportada.*PDF"):
        convert_file(sample("txt"), "xlsx", out_dir=tmp_path)


def test_same_format_is_rejected(sample, tmp_path):
    with pytest.raises(ConversionError, match="já está"):
        convert_file(sample("png"), "png", out_dir=tmp_path)


def test_failed_conversion_leaves_no_files(tmp_path):
    src = tmp_path / "quebrado.png"
    src.write_bytes(b"\x89PNG\r\n\x1a\n" + b"lixo" * 10)
    out = tmp_path / "out"
    with pytest.raises(ConversionError):
        convert_file(src, "jpg", out_dir=out)
    assert list(out.iterdir()) == []


def test_media_unavailable_without_ffmpeg(monkeypatch, tmp_path):
    media.ffmpeg_path.cache_clear()
    monkeypatch.setenv("PATH", "")
    monkeypatch.delenv("CONVERSOR_FFMPEG", raising=False)
    try:
        assert "mp3" not in registry.targets_for("wav")
        src = tmp_path / "a.wav"
        src.write_bytes(b"RIFF\x00\x00\x00\x00WAVEfmt ")
        with pytest.raises(ConversionError, match="componente ausente"):
            convert_file(src, "mp3", out_dir=tmp_path)
    finally:
        media.ffmpeg_path.cache_clear()


@pytest.mark.skipif(not HAS_FFMPEG, reason="ffmpeg não instalado")
def test_video_to_audio_has_duration(sample, tmp_path):
    import subprocess

    out = convert_file(sample("mp4"), "mp3", out_dir=tmp_path).output
    dur = subprocess.run(
        ["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", str(out)],
        capture_output=True, text=True, check=True,
    ).stdout
    assert 0.5 < float(dur) < 2.0


@pytest.mark.skipif(not HAS_FFMPEG, reason="ffmpeg não instalado")
def test_encoder_list_is_read_from_installed_ffmpeg():
    assert media.tem_codificador("flac") and media.tem_codificador("pcm_s16le")
    assert not media.tem_codificador("nao-existe")


@pytest.mark.skipif(not HAS_FFMPEG, reason="ffmpeg não instalado")
@pytest.mark.parametrize("canais", [1, 6])
def test_ogg_without_libvorbis_uses_internal_encoder_and_warns(tmp_path, monkeypatch, canais):
    """ffmpeg do Homebrew sem libvorbis: OGG sai com o Vorbis interno, em estéreo, com aviso."""
    import subprocess

    src = tmp_path / f"tom{canais}.wav"
    subprocess.run(["ffmpeg", "-nostdin", "-loglevel", "error", "-f", "lavfi",
                    "-i", "sine=frequency=440:duration=1", "-ac", str(canais), str(src)], check=True)
    real = media.tem_codificador
    monkeypatch.setattr(media, "tem_codificador", lambda n: n != "libvorbis" and real(n))

    r = convert_file(src, "ogg", out_dir=tmp_path)
    info = subprocess.run(
        ["ffprobe", "-v", "error", "-show_entries", "stream=codec_name,channels:format=duration",
         "-of", "csv=p=0", str(r.output)], capture_output=True, text=True, check=True,
    ).stdout.split()
    assert info[0] == "vorbis,2"
    assert 0.5 < float(info[1]) < 2.0
    assert any("libvorbis" in w for w in r.warnings)


@pytest.mark.parametrize(
    "text,expected",
    [
        ("a;b;c\n1;2,5;3\n4;5,5;6\n", ";"),
        ("a,b,c\n1,2.5,3\n", ","),
        ("a\tb\n1\t2\n", "\t"),
        ('a,b\n"x;y",2\n', ","),
    ],
)
def test_guess_delimiter(text, expected):
    from conversor.converters.data import guess_delimiter

    assert guess_delimiter(text) == expected
