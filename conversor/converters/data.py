"""Conversões de dados tabulares: CSV, JSON e XLSX.

Princípio: não alterar dados. Valores de CSV são lidos como texto e continuam
texto (CEP "01234" não vira 1234). Valores do XLSX mantêm o tipo da célula.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
from pathlib import Path

from conversor.core.registry import register

TABLE = ["csv", "json", "xlsx"]

Table = tuple[list[str], list[list[object]]]


def _warn(options: dict, msg: str) -> None:
    options.setdefault("_warnings", []).append(msg)


# ---------- leitura ----------

def _decode(src: Path, options: dict) -> str:
    raw = src.read_bytes()
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Arquivos salvos pelo Excel antigo no Windows costumam vir em cp1252.
        _warn(options, f"{src.name} não está em UTF-8; lido como Windows-1252.")
        return raw.decode("cp1252", errors="replace")


def guess_delimiter(text: str) -> str:
    """Escolhe o separador que gera o mesmo número de colunas em mais linhas.

    Não usamos csv.Sniffer: em CSV brasileiro ("10,50" com ";" separando)
    ele escolhe "," com frequência.
    """
    lines = [ln for ln in text.splitlines()[:50] if ln.strip()]
    best, best_score = ",", (0.0, 0)
    for delim in (";", ",", "\t", "|"):  # ordem = desempate
        counts = [len(r) for r in csv.reader(lines, delimiter=delim)]
        if not counts:
            continue
        mode = max(set(counts), key=counts.count)
        if mode < 2:
            continue
        score = (counts.count(mode) / len(counts), mode)
        if score > best_score:
            best, best_score = delim, score
    return best


def read_csv(src: Path, options: dict) -> Table:
    text = _decode(src, options)
    delim = options.get("csv_delimiter_in") or guess_delimiter(text)
    rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delim))
    rows = [r for r in rows if any(cell.strip() for cell in r)]
    if not rows:
        raise ValueError("CSV sem dados")
    return rows[0], rows[1:]


def read_json(src: Path, options: dict) -> Table:
    data = json.loads(_decode(src, options))
    if isinstance(data, dict):
        lists = [k for k, v in data.items() if isinstance(v, list)]
        if len(lists) == 1:
            _warn(options, f"Usando a lista da chave \"{lists[0]}\" do JSON.")
            data = data[lists[0]]
        else:
            raise ValueError(
                "O JSON precisa ser uma lista de objetos (ex.: [{\"a\": 1}, ...])."
            )
    if not isinstance(data, list) or not data:
        raise ValueError("O JSON precisa ser uma lista não vazia.")

    if all(isinstance(item, list) for item in data):
        return [str(h) for h in data[0]], [list(r) for r in data[1:]]
    if not all(isinstance(item, dict) for item in data):
        raise ValueError("Todos os itens do JSON devem ser objetos (ou todos listas).")

    header: list[str] = []
    for item in data:  # une as chaves de todos os objetos, na ordem em que aparecem
        for k in item:
            if k not in header:
                header.append(k)
    nested = False
    rows = []
    for item in data:
        row = []
        for k in header:
            v = item.get(k)
            if isinstance(v, (dict, list)):
                nested = True
                v = json.dumps(v, ensure_ascii=False)
            row.append(v)
        rows.append(row)
    if nested:
        _warn(options, "Valores aninhados (objetos/listas) foram gravados como texto JSON.")
    return header, rows


def read_xlsx(src: Path, options: dict) -> Table:
    from openpyxl import load_workbook

    wb = load_workbook(src, read_only=True, data_only=True)
    try:
        if len(wb.worksheets) > 1:
            _warn(
                options,
                f"{src.name} tem {len(wb.worksheets)} planilhas; apenas a primeira "
                f"(\"{wb.worksheets[0].title}\") foi convertida.",
            )
        rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        raise ValueError("Planilha vazia")
    header = ["" if h is None else str(h) for h in rows[0]]
    return header, rows[1:]


READERS = {"csv": read_csv, "json": read_json, "xlsx": read_xlsx}


# ---------- escrita ----------

def _json_value(v: object) -> object:
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return v.isoformat()
    return v


def _unique_keys(header: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for i, h in enumerate(header, start=1):
        key = h.strip() or f"coluna_{i}"
        if key in seen:
            seen[key] += 1
            key = f"{key}_{seen[key]}"
        else:
            seen[key] = 1
        out.append(key)
    return out


def write_csv(table: Table, dst: Path, options: dict) -> None:
    header, rows = table
    # ";" é o separador que o Excel em português (Brasil) espera ao abrir com duplo clique.
    delim = options.get("csv_delimiter_out", ";")
    # utf-8-sig (com BOM) faz o Excel reconhecer acentos corretamente.
    with open(dst, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=delim)
        w.writerow(header)
        for r in rows:
            w.writerow(["" if v is None else _json_value(v) for v in r])


def write_json(table: Table, dst: Path, options: dict) -> None:
    header, rows = table
    keys = _unique_keys(header)
    if keys != header:
        _warn(options, "Cabeçalhos vazios ou repetidos foram renomeados no JSON.")
    records = []
    for r in rows:
        r = list(r) + [None] * (len(keys) - len(r))
        records.append({k: _json_value(v) for k, v in zip(keys, r)})
    dst.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")


def write_xlsx(table: Table, dst: Path, options: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    header, rows = table
    wb = Workbook()
    ws = wb.active
    ws.title = "Dados"
    cleaned = False
    for r_idx, row in enumerate([header, *rows], start=1):
        for c_idx, v in enumerate(row, start=1):
            if isinstance(v, str) and ILLEGAL_CHARACTERS_RE.search(v):
                v = ILLEGAL_CHARACTERS_RE.sub("", v)
                cleaned = True
            cell = ws.cell(row=r_idx, column=c_idx, value=v)
            # Texto começando com "=" seria interpretado como fórmula pelo Excel.
            # Mantemos como texto para não alterar o dado (e evitar injeção de fórmula).
            if isinstance(v, str) and v.startswith("="):
                cell.data_type = "s"
    if cleaned:
        _warn(options, "Caracteres de controle inválidos no Excel foram removidos.")
    wb.save(dst)


WRITERS = {"csv": write_csv, "json": write_json, "xlsx": write_xlsx}


@register(TABLE, TABLE, category="Dados")
def convert_table(src: Path, dst: Path, options: dict) -> None:
    source = options["_source_format"]
    target = dst.suffix.lstrip(".")
    WRITERS[target](READERS[source](src, options), dst, options)
