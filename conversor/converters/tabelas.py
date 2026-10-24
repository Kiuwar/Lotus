"""Planilhas → tabela para ler e imprimir: XLSX, CSV e JSON → HTML, Markdown e PDF.

Todas as abas do XLSX entram (cada uma com o seu título). É uma conversão de
APRESENTAÇÃO: números do XLSX aparecem com vírgula decimal, porcentagens como
"15%" e datas como dd/mm/aaaa. CSV continua exatamente como está (é texto).
Cores, gráficos, larguras e células mescladas não são mantidos (aviso).

Fórmula sem valor salvo: o Lotus não calcula fórmulas. Um XLSX gravado por
outro programa sem o resultado calculado mostra a célula vazia, com aviso.
"""

from __future__ import annotations

import datetime as dt
import html
import re
from dataclasses import dataclass, field
from pathlib import Path

from conversor.converters.data import READERS, _warn
from conversor.core.engine import ConversionError
from conversor.core.registry import register

LIMITE_LINHAS_PDF = 5000  # por aba: um PDF maior fica lento e ilegível


@dataclass
class Aba:
    titulo: str
    linhas: list[list[str]] = field(default_factory=list)
    oculta: bool = False


# ---------- leitura ----------

def _numero(v: float | int, formato: str | None) -> str:
    formato = formato or "General"
    if "%" in formato:
        m = re.search(r"0\.?(0*)%", formato)
        casas = len(m.group(1)) if m else 0
        return f"{v * 100:.{casas}f}".replace(".", ",") + "%"
    if isinstance(v, float):
        if v.is_integer() and abs(v) < 1e15:
            texto = str(int(v))
        else:
            texto = repr(v).replace(".", ",")
    else:
        texto = str(v)
    if "R$" in formato:
        texto = "R$ " + texto
    return texto


def _celula(v: object, formato: str | None = None) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "VERDADEIRO" if v else "FALSO"
    if isinstance(v, dt.datetime):
        if (v.hour, v.minute, v.second) == (0, 0, 0):
            return v.strftime("%d/%m/%Y")
        return v.strftime("%d/%m/%Y %H:%M" if v.second == 0 else "%d/%m/%Y %H:%M:%S")
    if isinstance(v, dt.date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, dt.time):
        return v.strftime("%H:%M" if v.second == 0 else "%H:%M:%S")
    if isinstance(v, (int, float)):
        return _numero(v, formato)
    return str(v)


def _aparar(linhas: list[list[str]]) -> list[list[str]]:
    """Tira linhas vazias e colunas vazias à direita; iguala a largura."""
    linhas = [l for l in linhas if any(c.strip() for c in l)]
    largura = max((max((i + 1 for i, c in enumerate(l) if c.strip()), default=0) for l in linhas), default=0)
    return [(l + [""] * largura)[:largura] for l in linhas]


def ler_xlsx(src: Path, options: dict) -> list[Aba]:
    from openpyxl import load_workbook

    valores = load_workbook(src, read_only=True, data_only=True)
    formulas = load_workbook(src, read_only=True, data_only=False)
    abas: list[Aba] = []
    sem_valor = 0
    try:
        for ws, wf in zip(valores.worksheets, formulas.worksheets):
            linhas: list[list[str]] = []
            for rv, rf in zip(ws.iter_rows(), wf.iter_rows()):
                linha = []
                for cv, cf in zip(rv, rf):
                    valor = getattr(cv, "value", None)
                    if valor is None and isinstance(getattr(cf, "value", None), str) \
                            and cf.value.startswith("="):
                        sem_valor += 1
                    linha.append(_celula(valor, getattr(cv, "number_format", None)))
                linhas.append(linha)
            abas.append(Aba(ws.title, _aparar(linhas), oculta=ws.sheet_state != "visible"))
    finally:
        valores.close()
        formulas.close()
    if sem_valor:
        _warn(options, f"{sem_valor} fórmula(s) sem resultado salvo no arquivo aparecem vazias "
                       "(o Lotus não calcula fórmulas). Como resolver: abra no Excel, salve e "
                       "converta de novo.")
    ocultas = [a.titulo for a in abas if a.oculta]
    if ocultas:
        _warn(options, f"Abas ocultas também foram incluídas: {', '.join(ocultas)}.")
    _warn(options, "Cores, gráficos, larguras e células mescladas da planilha não são mantidos.")
    return abas


def ler_planilha(src: Path, origem: str, options: dict) -> list[Aba]:
    if origem == "xlsx":
        abas = ler_xlsx(src, options)
    else:
        cabecalho, linhas = READERS[origem](src, options)
        tabela = [[_celula(c) for c in cabecalho]] + [[_celula(c) for c in l] for l in linhas]
        abas = [Aba(src.stem, _aparar(tabela))]
    if not any(a.linhas for a in abas):
        raise ConversionError(f"{src.name} não tem dados para mostrar.")
    return abas


# ---------- escrita ----------

_CSS = """body{font-family:Segoe UI,-apple-system,Helvetica,Arial,sans-serif;margin:24px;color:#1d1d1f}
h1{font-size:20px}h2{font-size:16px;margin-top:28px}
table{border-collapse:collapse;font-size:13px}
th,td{border:1px solid #c7c7cc;padding:4px 8px;text-align:left;vertical-align:top}
th{background:#f2f2f7;position:sticky;top:0}
tbody tr:nth-child(even){background:#fafafa}
@media print{th{position:static}}"""


def para_html(abas: list[Aba], titulo: str) -> str:
    partes = []
    varias = len(abas) > 1
    for a in abas:
        if varias:
            partes.append(f"<h2>{html.escape(a.titulo)}{' (oculta)' if a.oculta else ''}</h2>")
        if not a.linhas:
            partes.append("<p><em>(aba vazia)</em></p>")
            continue
        cab = "".join(f"<th>{html.escape(c)}</th>" for c in a.linhas[0])
        corpo = "\n".join("<tr>" + "".join(f"<td>{html.escape(c)}</td>" for c in l) + "</tr>"
                          for l in a.linhas[1:])
        partes.append(f"<table>\n<thead><tr>{cab}</tr></thead>\n<tbody>\n{corpo}\n</tbody>\n</table>")
    return ("<!DOCTYPE html>\n<html lang=\"pt-BR\">\n<head>\n<meta charset=\"utf-8\">\n"
            f"<title>{html.escape(titulo)}</title>\n<style>\n{_CSS}\n</style>\n</head>\n<body>\n"
            f"<h1>{html.escape(titulo)}</h1>\n" + "\n".join(partes) + "\n</body>\n</html>\n")


def _md(c: str) -> str:
    return re.sub(r"([\\|`*_\[\]<>])", r"\\\1", c).replace("\r", "").replace("\n", "<br>")


def para_md(abas: list[Aba]) -> str:
    saida = []
    for a in abas:
        if len(abas) > 1:
            saida += [f"## {_md(a.titulo)}{' (oculta)' if a.oculta else ''}", ""]
        if not a.linhas:
            saida += ["*(aba vazia)*", ""]
            continue
        saida.append("| " + " | ".join(_md(c) for c in a.linhas[0]) + " |")
        saida.append("|" + "---|" * len(a.linhas[0]))
        saida += ["| " + " | ".join(_md(c) for c in l) + " |" for l in a.linhas[1:]]
        saida.append("")
    return "\n".join(saida).rstrip() + "\n"


def para_pdf(abas: list[Aba], dst: Path, titulo: str, options: dict) -> None:
    from fpdf import FPDF
    from fpdf.fonts import FontFace

    from conversor.converters.docx_escrita import _VARIANTES
    from conversor.converters.documents import _FONT_CANDIDATES

    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.set_title(titulo)
    pdf.set_creator("Lotus")
    fonte = next((p for p in _FONT_CANDIDATES if p.is_file()), None)
    if fonte:
        negrito = fonte.with_name(_VARIANTES.get(fonte.name, ("",))[0] or fonte.name)
        pdf.add_font("Texto", "", str(fonte))
        pdf.add_font("Texto", "B", str(negrito if negrito.is_file() else fonte))
        familia = "Texto"
    else:
        familia = "helvetica"

    def seguro(t: str) -> str:
        if fonte:
            return t
        s = t.encode("latin-1", errors="replace").decode("latin-1")
        if s != t:
            _warn(options, 'Nenhuma fonte Unicode encontrada; alguns caracteres viraram "?".')
        return s

    for a in abas:
        colunas = len(a.linhas[0]) if a.linhas else 0
        pdf.add_page(orientation="L" if colunas > 6 else "P")
        pdf.set_font(familia, "B", 13)
        nome = a.titulo + (" (oculta)" if a.oculta else "")
        pdf.cell(0, 8, seguro(nome if len(abas) > 1 else titulo), new_x="LMARGIN", new_y="NEXT")
        if not a.linhas:
            pdf.set_font(familia, "", 10)
            pdf.cell(0, 6, "(aba vazia)")
            continue
        linhas = a.linhas
        if len(linhas) - 1 > LIMITE_LINHAS_PDF:
            _warn(options, f"A aba \"{a.titulo}\" tem {len(linhas) - 1} linhas; o PDF mostra as "
                           f"primeiras {LIMITE_LINHAS_PDF}. Para ver tudo, converta para HTML.")
            linhas = linhas[: LIMITE_LINHAS_PDF + 1]
        if colunas > 30:
            _warn(options, f"A aba \"{a.titulo}\" tem {colunas} colunas; no PDF elas ficam muito "
                           "estreitas. Para ler melhor, converta para HTML.")
        pdf.set_font(familia, "", 9 if colunas <= 8 else 7 if colunas <= 16 else 5)
        cinza = FontFace(emphasis="BOLD", fill_color=(242, 242, 247))
        with pdf.table(headings_style=cinza, line_height=pdf.font_size * 1.6, text_align="LEFT",
                       repeat_headings=1, padding=1) as tabela:
            for linha in linhas:
                r = tabela.row()
                for c in linha:
                    r.cell(seguro(c))
    pdf.output(str(dst))


NOTA = "Todas as abas; números com vírgula decimal e datas dd/mm/aaaa; sem cores nem gráficos."


@register(["xlsx", "csv", "json"], ["html", "md", "pdf"], category="Dados", notes=NOTA)
def planilha_para_tabela(src: Path, dst: Path, options: dict) -> None:
    origem = options.get("_source_format") or src.suffix.lstrip(".").lower()
    abas = ler_planilha(src, origem, options)
    alvo = dst.suffix.lstrip(".")
    if alvo == "html":
        dst.write_text(para_html(abas, src.stem), encoding="utf-8")
    elif alvo == "md":
        dst.write_text(para_md(abas), encoding="utf-8")
    else:
        para_pdf(abas, dst, src.stem, options)
